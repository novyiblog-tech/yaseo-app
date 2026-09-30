"""Сборка финмодели yaSEO в xlsx.

Запуск: /opt/homebrew/bin/python3 finmodel/build.py
Результат: finmodel/yaseo-finmodel.xlsx — все расчёты формулами, вводные на листе «Параметры».
Значения вводных и их источники лежат в params.py.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from params import MONTHS, PARAMS, SCENARIOS, TARIFFS

OUT = Path(__file__).parent / "yaseo-finmodel.xlsx"
PSHEET = "Параметры"

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
INPUT = Font(name="Arial", size=10, color="0000FF")
LINK = Font(name="Arial", size=10, color="008000")
KEY = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="E7E6E6")

RUB = '#,##0;(#,##0);"-"'
RUB2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0.0;(#,##0.0);"-"'
FORMATS = {"rub": RUB, "rub2": RUB2, "pct": PCT, "num": NUM}

SC_COL = {name: get_column_letter(2 + i) for i, name in enumerate(SCENARIOS)}
BASE = SCENARIOS[1]


def build_params(wb):
    ws = wb.active
    ws.title = PSHEET
    ws["A1"] = "yaSEO — параметры финмодели"
    ws["A1"].font = TITLE
    ws["A2"] = ("Синие числа — вводные, их можно менять. Жёлтая заливка — допущения без замера, "
                "их проверяет бета. Чёрные и зелёные ячейки — формулы, их не трогать. "
                "Общий параметр меняется в столбце «" + BASE + "», остальные два подтянутся.")
    ws["A2"].font = FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 44

    head = ["Параметр", *SCENARIOS, "Единица", "Источник и пояснение"]
    for c, text in enumerate(head, 1):
        cell = ws.cell(row=4, column=c, value=text)
        cell.font, cell.fill = BOLD, HEAD

    rows = {}
    r = 5
    for item in PARAMS:
        if "section" in item:
            ws.cell(row=r, column=1, value=item["section"]).font = BOLD
            r += 1
            continue
        rows[item["key"]] = r
        ws.cell(row=r, column=1, value=item["name"]).font = FONT
        fmt = FORMATS[item["fmt"]]
        values = item["values"]
        for i, name in enumerate(SCENARIOS):
            cell = ws.cell(row=r, column=2 + i)
            if isinstance(values, (list, tuple)):
                cell.value, cell.font = values[i], INPUT
            elif name == BASE:
                cell.value, cell.font = values, INPUT
            else:
                cell.value, cell.font = f"={SC_COL[BASE]}{r}", FONT
            cell.number_format = fmt
            if item.get("assumption"):
                cell.fill = KEY
        ws.cell(row=r, column=5, value=item["unit"]).font = FONT
        src = ws.cell(row=r, column=6, value=item["source"])
        src.font = FONT
        src.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    ws.column_dimensions["A"].width = 52
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 90
    ws.freeze_panes = "B5"
    return rows


def build_scenario(wb, name, prow):
    ws = wb.create_sheet(name)
    col = SC_COL[name]

    def p(key):
        return f"'{PSHEET}'!${col}${prow[key]}"

    ws["A1"] = f"yaSEO — сценарий «{name.lower()}», {MONTHS} месяцев, ₽"
    ws["A1"].font = TITLE
    ws.cell(row=3, column=1, value="Месяц").font = BOLD
    ws.cell(row=3, column=1).fill = HEAD
    first, last = 2, 1 + MONTHS
    for m in range(1, MONTHS + 1):
        cell = ws.cell(row=3, column=1 + m, value=m)
        cell.font, cell.fill = BOLD, HEAD

    layout = [
        ("Поток", None, None),
        ("org", "Посетители бесплатных каналов", "num"),
        ("paid", "Посетители из рекламы", "num"),
        ("vis", "Посетители всего", "num"),
        ("reg", "Бесплатные проверки (регистрации)", "num"),
        ("new", "Новые подписчики", "num"),
        ("one", "Разовые аудиты", "num"),
        ("Подписчики", None, None),
        ("churned", "Ушло подписчиков", "num"),
        ("act", "Активные подписчики", "num"),
        *[(f"act_{t}", f"  из них {TARIFFS[t]}", "num") for t in TARIFFS],
        ("Выручка", None, None),
        ("rev_sub", "Подписки", "rub"),
        ("rev_one", "Разовые аудиты", "rub"),
        ("rev", "Выручка всего", "rub"),
        ("Переменные расходы", None, None),
        ("api_sub", "API Яндекса: подписчики", "rub"),
        ("api_one", "API Яндекса: разовые аудиты", "rub"),
        ("api_free", "API Яндекса: бесплатные проверки", "rub"),
        ("llm", "LLM: тексты рекомендаций", "rub"),
        ("acq", "Эквайринг", "rub"),
        ("Налоги и взносы", None, None),
        ("rev12", "Выручка за 12 месяцев (для порога НДС)", "rub"),
        ("tax", "Налог УСН", "rub"),
        ("vat", "НДС", "rub"),
        ("ins", "Страховые взносы ИП", "rub"),
        ("Постоянные расходы", None, None),
        ("fixed", "Сервер, база, касса, домен, прочее", "rub"),
        ("mkt", "Реклама", "rub"),
        ("labor", "Работа владельцев", "rub"),
        ("Итог", None, None),
        ("cost", "Расходы всего", "rub"),
        ("profit", "Прибыль за месяц", "rub"),
        ("cum", "Накопленный итог", "rub"),
        ("f_profit", "Месяц прибыльный (1 — да)", "num"),
        ("f_cum", "Вложения возвращены (1 — да)", "num"),
    ]
    row = {}
    r = 4
    for key, label, fmt in layout:
        if label is None:
            ws.cell(row=r, column=1, value=key).font = BOLD
        else:
            row[key] = r
            ws.cell(row=r, column=1, value=label).font = FONT
        r += 1

    tkeys = list(TARIFFS)
    fixed_keys = ["fix_server", "fix_db", "fix_kassa", "fix_domain", "fix_other"]

    for m in range(1, MONTHS + 1):
        c = get_column_letter(1 + m)
        prev = get_column_letter(m)

        def ref(key, column=c):
            return f"{column}{row[key]}"

        f = {}
        f["org"] = f"={p('org_start')}" if m == 1 else f"={ref('org', prev)}*(1+{p('org_growth')})"
        f["paid"] = f"=IF({p('cpv')}>0,{p('mkt_budget')}/{p('cpv')},0)"
        f["vis"] = f"={ref('org')}+{ref('paid')}"
        f["reg"] = f"={ref('vis')}*{p('conv_reg')}"
        f["new"] = f"={ref('reg')}*{p('conv_paid')}"
        f["one"] = f"={ref('reg')}*{p('conv_one')}"
        f["churned"] = "=0" if m == 1 else f"={ref('act', prev)}*{p('churn')}"
        f["act"] = (f"={ref('new')}" if m == 1
                    else f"={ref('act', prev)}+{ref('new')}-{ref('churned')}")
        for t in tkeys:
            f[f"act_{t}"] = f"={ref('act')}*{p('mix_' + t)}"
        f["rev_sub"] = "=" + "+".join(f"{ref('act_' + t)}*{p('price_' + t)}" for t in tkeys)
        f["rev_one"] = f"={ref('one')}*{p('price_one')}"
        f["rev"] = f"={ref('rev_sub')}+{ref('rev_one')}"
        f["api_sub"] = ("=(" + "+".join(f"{ref('act_' + t)}*{p('api_' + t)}" for t in tkeys)
                        + f")*{p('usage')}")
        f["api_one"] = f"={ref('one')}*{p('api_audit')}"
        f["api_free"] = f"={ref('reg')}*{p('api_express')}"
        f["llm"] = ("=(" + "+".join(f"{ref('act_' + t)}*{p('audits_' + t)}" for t in tkeys)
                    + f")*{p('usage')}*{p('llm_report')}+{ref('one')}*{p('llm_report')}"
                    + f"+{ref('reg')}*{p('llm_express')}")
        f["acq"] = f"={ref('rev')}*{p('acq_rate')}"
        start = get_column_letter(max(first, 1 + m - 11))
        f["rev12"] = f"=SUM({start}{row['rev']}:{c}{row['rev']})"
        f["tax"] = f"={ref('rev')}*{p('tax_rate')}"
        f["vat"] = f"=IF({ref('rev12')}>{p('vat_threshold')},{ref('rev')}*{p('vat_rate')},0)"
        f["ins"] = (f"={p('ins_fixed')}/12+MAX(0,{ref('rev')}-{p('ins_base')}/12)"
                    f"*{p('ins_rate')}")
        f["fixed"] = "=" + "+".join(p(k) for k in fixed_keys)
        f["mkt"] = f"={p('mkt_budget')}"
        f["labor"] = f"={p('labor')}"
        cost_keys = ["api_sub", "api_one", "api_free", "llm", "acq", "tax", "vat", "ins",
                     "fixed", "mkt", "labor"]
        f["cost"] = "=" + "+".join(ref(k) for k in cost_keys)
        f["profit"] = f"={ref('rev')}-{ref('cost')}"
        f["cum"] = (f"={ref('profit')}" if m == 1
                    else f"={ref('cum', prev)}+{ref('profit')}")
        f["f_profit"] = f"=IF({ref('profit')}>=0,1,0)"
        f["f_cum"] = f"=IF({ref('cum')}>=0,1,0)"

        fmts = {key: fmt for key, label, fmt in layout if label}
        for key, formula in f.items():
            cell = ws.cell(row=row[key], column=1 + m, value=formula)
            cell.number_format = FORMATS[fmts[key]]
            cell.font = LINK if "'" + PSHEET + "'" in formula else FONT

    for key in ("rev", "cost", "profit", "cum", "act"):
        ws.cell(row=row[key], column=1).font = BOLD

    ws.column_dimensions["A"].width = 44
    for m in range(1, MONTHS + 1):
        ws.column_dimensions[get_column_letter(1 + m)].width = 12
    ws.freeze_panes = "B4"
    return row


def build_summary(wb, prow, srows):
    ws = wb.create_sheet("Сводка", 0)
    ws["A1"] = "yaSEO — сводка по трём сценариям"
    ws["A1"].font = TITLE
    ws["A2"] = ("Все числа считаются из листа «Параметры». Это модель на допущениях: конверсии, "
                "отток и поток посетителей до беты не измерены.")
    ws["A2"].font = FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 32

    for c, text in enumerate(["Показатель", *SCENARIOS], 1):
        cell = ws.cell(row=4, column=c, value=text)
        cell.font, cell.fill = BOLD, HEAD

    c1, c12, c13, c24 = "B", get_column_letter(13), get_column_letter(14), get_column_letter(25)
    rng_all = f"{c1}{{r}}:{c24}{{r}}"
    lines = [
        ("Активные подписчики, месяц 12", "num", lambda s, R: f"='{s}'!{c12}{R['act']}"),
        ("Активные подписчики, месяц 24", "num", lambda s, R: f"='{s}'!{c24}{R['act']}"),
        ("Выручка подписок за месяц 12", "rub", lambda s, R: f"='{s}'!{c12}{R['rev_sub']}"),
        ("Выручка подписок за месяц 24", "rub", lambda s, R: f"='{s}'!{c24}{R['rev_sub']}"),
        ("Выручка за первый год", "rub",
         lambda s, R: f"=SUM('{s}'!{c1}{R['rev']}:{c12}{R['rev']})"),
        ("Выручка за второй год", "rub",
         lambda s, R: f"=SUM('{s}'!{c13}{R['rev']}:{c24}{R['rev']})"),
        ("Прибыль за первый год", "rub",
         lambda s, R: f"=SUM('{s}'!{c1}{R['profit']}:{c12}{R['profit']})"),
        ("Прибыль за второй год", "rub",
         lambda s, R: f"=SUM('{s}'!{c13}{R['profit']}:{c24}{R['profit']})"),
        ("Накопленный итог на конец месяца 24", "rub",
         lambda s, R: f"='{s}'!{c24}{R['cum']}"),
        ("Самая глубокая яма (сколько денег нужно вложить)", "rub",
         lambda s, R: "=MIN(0,MIN('" + s + "'!" + rng_all.format(r=R['cum']) + "))"),
        ("Первый прибыльный месяц", "text",
         lambda s, R: "=IFERROR(MATCH(1,'" + s + "'!" + rng_all.format(r=R['f_profit'])
         + ',0),"за 24 месяца нет")'),
        ("Месяц возврата вложений", "text",
         lambda s, R: "=IFERROR(MATCH(1,'" + s + "'!" + rng_all.format(r=R['f_cum'])
         + ',0),"за 24 месяца нет")'),
        ("Подписчиков для выхода в ноль при текущих расходах", "num", None),
        ("Средний чек подписки", "rub", None),
        ("Маржа с подписчика в месяц", "rub", None),
        ("Срок жизни подписчика, месяцев", "num", None),
        ("Доход с подписчика за срок жизни (LTV)", "rub", None),
        ("Цена привлечения подписчика из рекламы (CAC)", "rub", None),
    ]
    tkeys = list(TARIFFS)
    pos = {label: 5 + i for i, (label, _, _) in enumerate(lines)}
    r = 5
    for label, fmt, fn in lines:
        ws.cell(row=r, column=1, value=label).font = FONT
        for i, s in enumerate(SCENARIOS):
            col = SC_COL[s]
            cell = ws.cell(row=r, column=2 + i)
            cell.font = LINK

            def p(key, col=col):
                return f"'{PSHEET}'!${col}${prow[key]}"

            if fn is not None:
                cell.value = fn(s, srows[s])
            elif label.startswith("Средний чек"):
                cell.value = "=" + "+".join(f"{p('mix_' + t)}*{p('price_' + t)}" for t in tkeys)
            elif label.startswith("Маржа"):
                api = "+".join(f"{p('mix_' + t)}*{p('api_' + t)}" for t in tkeys)
                aud = "+".join(f"{p('mix_' + t)}*{p('audits_' + t)}" for t in tkeys)
                check = f"{get_column_letter(2 + i)}{pos['Средний чек подписки']}"
                cell.value = (f"={check}*(1-{p('acq_rate')}-{p('tax_rate')}-{p('ins_rate')})"
                              f"-({api})*{p('usage')}-({aud})*{p('usage')}*{p('llm_report')}")
                cell.font = FONT
            elif label.startswith("Срок жизни"):
                cell.value = f"=IF({p('churn')}>0,1/{p('churn')},0)"
            elif label.startswith("Доход с подписчика"):
                cl = get_column_letter(2 + i)
                cell.value = (f"={cl}{pos['Маржа с подписчика в месяц']}"
                              f"*{cl}{pos['Срок жизни подписчика, месяцев']}")
                cell.font = FONT
            elif label.startswith("Цена привлечения"):
                denom = f"{p('conv_reg')}*{p('conv_paid')}"
                cell.value = (f'=IF({p("mkt_budget")}=0,"рекламы нет",'
                              f"IF({denom}>0,{p('cpv')}/({denom}),0))")
            elif label.startswith("Подписчиков для выхода"):
                cl = get_column_letter(2 + i)
                fixed = "+".join(p(k) for k in ["fix_server", "fix_db", "fix_kassa",
                                                "fix_domain", "fix_other", "mkt_budget",
                                                "labor"])
                margin = f"{cl}{pos['Маржа с подписчика в месяц']}"
                cell.value = (f"=IF({margin}>0,({fixed}+{p('ins_fixed')}/12)/{margin},0)")
            cell.number_format = FORMATS.get(fmt, "0")
        r += 1

    ws.column_dimensions["A"].width = 56
    for col in "BCD":
        ws.column_dimensions[col].width = 22
    ws.freeze_panes = "B5"


def main():
    wb = Workbook()
    prow = build_params(wb)
    srows = {name: build_scenario(wb, name, prow) for name in SCENARIOS}
    build_summary(wb, prow, srows)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
